from __future__ import annotations

from pathlib import Path

from docx import Document

from .errors import fail


def describe_template(path: Path) -> dict:
    """Return a compact, deterministic description that can be sent to the VLM."""
    try:
        document = Document(path)
    except Exception as exc:
        raise fail("HF0203", "模板无法读取或不是有效DOCX") from exc
    if not document.tables:
        raise fail("HF0203")
    tables = []
    for table_index, table in enumerate(document.tables):
        seen: dict[object, tuple[int, int]] = {}
        cells = []
        for row_index, row in enumerate(table.rows):
            for col_index, cell in enumerate(row.cells):
                tc = cell._tc
                anchor = seen.get(tc)
                if anchor is None:
                    anchor = (row_index, col_index)
                    seen[tc] = anchor
                cells.append(
                    {
                        "row": row_index,
                        "col": col_index,
                        "anchorRow": anchor[0],
                        "anchorCol": anchor[1],
                        "merged": anchor != (row_index, col_index),
                        "fixedText": cell.text.strip() if anchor == (row_index, col_index) else "",
                    }
                )
        tables.append(
            {
                "tableIndex": table_index,
                "rows": len(table.rows),
                "cols": max((len(row.cells) for row in table.rows), default=0),
                "cells": cells,
            }
        )
    return {"tableCount": len(tables), "tables": tables}
