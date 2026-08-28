from __future__ import annotations

from pathlib import Path

from docx import Document

from .errors import fail
from .models import RecognizedTable


def _is_merged_duplicate(cell, seen: set[object]) -> bool:
    tc = cell._tc
    if tc in seen:
        return True
    seen.add(tc)
    return False


def _replace_text_preserving_format(cell, text: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        for run in extra.runs:
            run.text = ""


def fill_template(
    template_path: Path,
    output_path: Path,
    recognized_tables: list[RecognizedTable],
    preserve_original_text: bool,
) -> int:
    try:
        document = Document(template_path)
    except Exception as exc:
        raise fail("HF0203", "模板无法读取") from exc
    if not document.tables:
        raise fail("HF0203")
    if len(document.tables) != len(recognized_tables):
        raise fail("HF0204", f"识别到 {len(recognized_tables)} 个表格，模板包含 {len(document.tables)} 个表格")
    filled = 0
    for table, recognized in zip(document.tables, recognized_tables):
        row_count = len(table.rows)
        col_count = max((len(row.cells) for row in table.rows), default=0)
        if (row_count, col_count) != (recognized.rows, recognized.cols):
            raise fail(
                "HF0204",
                f"表格结构为 {recognized.rows}x{recognized.cols}，模板结构为 {row_count}x{col_count}",
            )
        values = {(cell.row, cell.col): cell.text for cell in recognized.cells}
        seen: set[object] = set()
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                if _is_merged_duplicate(cell, seen):
                    continue
                text = values.get((r, c), "").strip()
                if not text:
                    continue
                if preserve_original_text and cell.text.strip():
                    continue
                _replace_text_preserving_format(cell, text)
                filled += 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return filled
