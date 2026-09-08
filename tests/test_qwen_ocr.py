import json
from unittest.mock import patch

import numpy as np

from handwritten_form_api.ocr import QwenVlDocumentAnalyzer, _json_object


class FakeResponse:
    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def read(self):
        content = {
            "structureMatched": True,
            "mismatchReason": "",
            "tables": [{
                "tableIndex": 0, "rows": 1, "cols": 2,
                "cells": [{"row": 0, "col": 1, "text": "张三", "confidence": 0.91}],
            }],
        }
        return json.dumps({"choices": [{"message": {"content": json.dumps(content, ensure_ascii=False)}}]}).encode()


def test_qwen_receives_template_and_all_images_once():
    schema = {"tableCount": 1, "tables": [{
        "tableIndex": 0, "rows": 1, "cols": 2,
        "cells": [
            {"row": 0, "col": 0, "anchorRow": 0, "anchorCol": 0, "merged": False, "fixedText": "姓名"},
            {"row": 0, "col": 1, "anchorRow": 0, "anchorCol": 1, "merged": False, "fixedText": ""},
        ],
    }]}
    analyzer = QwenVlDocumentAnalyzer("http://model/v1", "local", "qwen", retries=0)
    images = [np.full((40, 120, 3), 255, dtype=np.uint8), np.full((40, 120, 3), 255, dtype=np.uint8)]
    with patch("urllib.request.urlopen", return_value=FakeResponse()) as called:
        tables = analyzer.analyze(images, schema, "zh-CN")
    assert tables[0].cells[0].text == "张三"
    assert called.call_count == 1
    body = json.loads(called.call_args.args[0].data.decode("utf-8"))
    assert body["response_format"] == {"type": "json_object"}
    image_parts = [part for part in body["messages"][1]["content"] if part["type"] == "image_url"]
    assert len(image_parts) == 2
    prompt = body["messages"][1]["content"][-1]["text"]
    assert '"row":0,"col":0,"fixedText":"姓名"' in prompt
    assert '"merged"' not in prompt
    assert "空白单元格必须省略" in prompt


def test_extracts_last_complete_json_from_ollama_reasoning():
    reasoning = '先参考示例 {"example": true}，最终结果：' \
        '{"structureMatched":true,"tables":[]}'
    extracted = json.loads(_json_object(reasoning))
    assert extracted == {"structureMatched": True, "tables": []}


def test_template_dimensions_override_model_dimensions():
    schema = {"tables": [{"tableIndex": 0, "rows": 11, "cols": 10,
                           "cells": [{"row": 1, "col": 2, "merged": False}]}]}
    data = {"tables": [{"tableIndex": 0, "rows": 2, "cols": 3,
                        "cells": [{"row": 1, "col": 2, "text": "resultFileId"}]}]}
    table = QwenVlDocumentAnalyzer._parse_tables(data, schema)[0]
    assert (table.rows, table.cols) == (11, 10)
    assert table.cells[0].text == "resultFileId"


def test_missing_tables_and_split_fragments_keep_template_structure():
    schema = {"tables": [
        {"tableIndex": i, "rows": 1, "cols": 2,
         "cells": [{"row": 0, "col": c, "merged": False} for c in range(2)]}
        for i in range(2)
    ]}
    data = {"tables": [
        {"tableIndex": 0, "cells": [{"row": 0, "col": 0, "text": "A"}]},
        {"tableIndex": 0, "cells": [{"row": 0, "col": 1, "text": "B"}]}
    ]}
    tables = QwenVlDocumentAnalyzer._parse_tables(data, schema)
    assert len(tables) == 2
    assert [cell.text for cell in tables[0].cells] == ["A", "B"]
    assert tables[1].cells == ()


def test_invalid_coordinates_and_conflicting_content_are_rejected():
    import pytest
    from handwritten_form_api.ocr import StructureMismatch

    schema = {"tables": [{"tableIndex": 0, "rows": 1, "cols": 1,
                           "cells": [{"row": 0, "col": 0, "merged": False}]}]}
    for fragments in [
        [{"tableIndex": 9, "cells": []}],
        [{"tableIndex": 0, "cells": [{"row": 1, "col": 0, "text": "X"}]}],
        [{"tableIndex": 0, "cells": [{"row": 0, "col": 0, "text": text}]}
         for text in ["A", "B"]],
    ]:
        with pytest.raises(StructureMismatch):
            QwenVlDocumentAnalyzer._parse_tables({"tables": fragments}, schema)


def test_new_response_produces_editable_word(tmp_path):
    from docx import Document
    from handwritten_form_api.pipeline import ConversionPipeline

    template = tmp_path / "template.docx"
    source = tmp_path / "source.png"
    output = tmp_path / "output.docx"
    document = Document()
    table = document.add_table(rows=11, cols=10)
    table.cell(0, 0).text = "序号"
    document.save(template)
    from PIL import Image
    Image.new("RGB", (100, 100), "white").save(source)
    response = {"choices": [{"message": {"content": json.dumps({"tables": [{
        "tableIndex": 0, "cells": [{"row": 1, "col": 2,
        "text": "resultFileId", "confidence": 0.95}]}]})}}]}
    analyzer = QwenVlDocumentAnalyzer("http://model/v1", "local", "qwen", retries=0)
    with patch.object(analyzer, "_request_with_retry", return_value=response) as request:
        result = ConversionPipeline(analyzer).run([source], template, output, "zh-CN", True)
    prompt = request.call_args.args[0]["messages"][1]["content"][-1]["text"]
    assert '"rows":11,"cols":10' in prompt
    assert '"rows":2,"cols":3' not in prompt
    assert "印刷文字" in prompt
    assert result["filledCellCount"] == 1
    reopened = Document(output)
    assert len(reopened.tables[0].rows) == 11
    assert reopened.tables[0].cell(0, 0).text == "序号"
    assert reopened.tables[0].cell(1, 2).text == "resultFileId"
    reopened.tables[0].cell(1, 2).text = "edited"
    reopened.save(output)
    assert Document(output).tables[0].cell(1, 2).text == "edited"
