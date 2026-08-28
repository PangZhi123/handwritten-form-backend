from pathlib import Path

from docx import Document

from handwritten_form_api.template_schema import describe_template


def test_template_schema_contains_fixed_and_writable_cells(tmp_path: Path):
    path = tmp_path / "template.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    document.save(path)
    schema = describe_template(path)
    assert schema["tableCount"] == 1
    assert schema["tables"][0]["cells"][0]["fixedText"] == "姓名"
    assert schema["tables"][0]["cells"][1]["fixedText"] == ""

