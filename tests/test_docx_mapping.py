from pathlib import Path

import pytest
from docx import Document

from handwritten_form_api.docx_mapping import fill_template
from handwritten_form_api.errors import BusinessError
from handwritten_form_api.models import RecognizedCell, RecognizedTable


def make_template(path: Path, rows: int = 2, cols: int = 2) -> None:
    document = Document()
    document.add_table(rows=rows, cols=cols)
    document.save(path)


def test_fill_template_creates_editable_copy(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "result.docx"
    make_template(template)
    recognized = [
        RecognizedTable(
            0,
            0,
            2,
            2,
            (
                RecognizedCell(0, 0, "姓名", 0.99),
                RecognizedCell(0, 1, "张三", 0.92),
                RecognizedCell(1, 0, "日期", 0.98),
                RecognizedCell(1, 1, "2026-08-28", 0.88),
            ),
        )
    ]
    assert fill_template(template, output, recognized, False) == 4
    result = Document(output)
    assert result.tables[0].cell(0, 1).text == "张三"
    assert Document(template).tables[0].cell(0, 1).text == ""


def test_structure_mismatch_is_rejected(tmp_path: Path):
    template = tmp_path / "template.docx"
    make_template(template)
    recognized = [RecognizedTable(0, 0, 3, 2, ())]
    with pytest.raises(BusinessError) as raised:
        fill_template(template, tmp_path / "result.docx", recognized, True)
    assert raised.value.code == "HF0204"


def test_preserve_original_text_skips_fixed_cell(tmp_path: Path):
    template = tmp_path / "template.docx"
    output = tmp_path / "result.docx"
    make_template(template, 1, 2)
    document = Document(template)
    document.tables[0].cell(0, 0).text = "固定标题"
    document.save(template)
    recognized = [RecognizedTable(0, 0, 1, 2, (
        RecognizedCell(0, 0, "不应覆盖", 0.9),
        RecognizedCell(0, 1, "填写内容", 0.9),
    ))]
    assert fill_template(template, output, recognized, True) == 1
    result = Document(output)
    assert result.tables[0].cell(0, 0).text == "固定标题"
    assert result.tables[0].cell(0, 1).text == "填写内容"

