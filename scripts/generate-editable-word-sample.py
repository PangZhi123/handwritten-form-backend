from __future__ import annotations

import argparse
import sys
import zipfile
from pathlib import Path

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_ROOT = PROJECT_ROOT / "src"
sys.path.insert(0, str(SRC_ROOT))

from handwritten_form_api.docx_mapping import fill_template  # noqa: E402
from handwritten_form_api.models import RecognizedCell, RecognizedTable  # noqa: E402


SAMPLE_VALUES = {
    (3, 10): "2024.07.08 20时30分",
    (4, 2): "2024.07.13 06时00分",
    (4, 10): "2024.07.13 08时30分",
    (7, 5): "1535",
    (8, 5): "-1",
    (9, 5): "220.5",
    (10, 5): "未入岩",
    (11, 5): "0.3‰",
    (12, 5): "220",
    (13, 5): "1.09",
    (14, 5): "32",
    (15, 5): "0.5%",
    (16, 5): "50",
}


def main() -> int:
    parser = argparse.ArgumentParser(description="不调用OCR，直接验证可编辑Word生成链路")
    parser.add_argument(
        "--template",
        type=Path,
        default=PROJECT_ROOT / "test-data" / "case-01" / "table_template.docx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=PROJECT_ROOT / "runtime" / "manual-test" / "editable-sample.docx",
    )
    args = parser.parse_args()

    template = args.template.resolve()
    output = args.output.resolve()
    source = Document(template)
    if len(source.tables) != 1:
        raise RuntimeError(f"样例脚本预期1个表格，实际为{len(source.tables)}个")
    table = source.tables[0]
    rows = len(table.rows)
    cols = max(len(row.cells) for row in table.rows)
    cells = tuple(
        RecognizedCell(row, col, text, 1.0)
        for (row, col), text in SAMPLE_VALUES.items()
    )
    filled = fill_template(
        template,
        output,
        [RecognizedTable(0, 0, rows, cols, cells)],
        preserve_original_text=True,
    )

    if not zipfile.is_zipfile(output):
        raise RuntimeError("输出不是有效的DOCX压缩包")
    verified = Document(output)
    if verified.tables[0].cell(7, 5).text != "1535":
        raise RuntimeError("输出Word中的测试值校验失败")

    print(f"生成成功: {output}")
    print(f"表格数量: {len(verified.tables)}")
    print(f"填充单元格: {filled}")
    print("文件结构: 有效DOCX，可用Microsoft Word打开并编辑")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
